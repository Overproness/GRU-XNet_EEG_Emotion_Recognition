"""Reconcile first-party identity/timing sources while keeping outcomes closed.

Agreement among exported metadata is distinguished from an original execution
log. Catalogue lengths are never promoted to observed playback endpoints.
"""
from collections import Counter
import io
import json
import math
from pathlib import Path
import re

from analyze_emo_mc_annotations import (numeric_identity, pair_trials,
                                      separate_initial_state_markers, trigger_mappings,
                                      is_subsequence)
from collect_emo_mc_historical_events import ROOT, PUBLIC
from qualify_emo_mc import REPO, PRIVATE, OLD, declared_orders, catalog, sha, write_json


def archive_orders(text):
    orders = declared_orders(text)
    # The alternate README explicitly documents distinct orders for sub-54.
    for context, word in [('ima', 'Imagery'), ('vid', 'Video')]:
        match = re.search(r'\*\*sub54 ' + word + r' sequence\*\*:.*?`(\[[^`]+\])`', text, re.S)
        if not match:
            raise ValueError('Missing explicit alternate-archive sub-54 order')
        import ast
        values = ast.literal_eval(match[1])
        if len(values) != 21 or len(set(values)) != 21:
            raise ValueError('Invalid alternate order')
        orders.setdefault('sub-54', {})[context] = values
    return orders


def clock_blocks(rows):
    """Retain source row order; backward clocks identify candidate blocks only."""
    result = []
    block = 1
    previous = -1
    for row in rows:
        onset = row['onset_s']
        if not math.isfinite(onset) or onset < 0 or row['duration_s'] != 0:
            raise ValueError('Invalid historical timing')
        if onset < previous:
            block += 1
        result.append(dict(onset_s=onset, description=f"TypeID: {row['trigger_code']}",
                           run_path=f'historical-clock-block-{block:02d}'))
        previous = onset
    return result


def compare_events(historical, authenticated, tolerance=0.000050001):
    if len(historical) != len(authenticated):
        return dict(ordered_markers_match_at_edf_precision=False, historical_rows=len(historical),
                    authenticated_rows=len(authenticated))
    matches = [h['trigger_code'] == int(re.fullmatch(r'TypeID:\s*(\d+)', a['description'])[1])
               and abs(h['onset_s'] - a['onset_s']) <= tolerance
               for h, a in zip(historical, authenticated)]
    return dict(ordered_markers_match_at_edf_precision=all(matches), historical_rows=len(historical),
                authenticated_rows=len(authenticated), matching_rows=sum(matches),
                tolerance_s=tolerance,
                precision_reason='EDF pilot TAL onsets have four decimal places; half a 0.0001s serialization step plus numerical roundoff.',
                maximum_onset_difference_s=max(abs(h['onset_s']-a['onset_s']) for h,a in zip(historical,authenticated)),
                run_ids_in_historical_table=False)


def video_endpoint(start, rating, length):
    """Return transparent timing diagnostics; neither proxy certifies playback."""
    if not all(math.isfinite(x) for x in (start, rating, length)):
        raise ValueError('Nonfinite endpoint')
    if length < 30 or rating <= start:
        raise ValueError('Invalid endpoint interval')
    catalog_end = start + length
    rating_guard_end = rating - 1
    overlap = max(0, min(catalog_end, rating_guard_end) - max(catalog_end-30, rating_guard_end-30))
    return dict(catalogue_duration_s=length, catalogue_end_candidate_s=catalog_end,
                rating_minus_1s_candidate_s=rating_guard_end,
                candidate_end_difference_s=rating_guard_end-catalog_end,
                overlap_between_two_30s_candidates_s=overlap,
                actual_playback_end_verified=False, fitting_allowed=False)


def protected_mask(context, symbol, number_map, reserved_keys):
    key = number_map.get((context, symbol))
    if key is None:
        return dict(material_key=None, status='unknown_identity_exclude', fitting_allowed=False)
    return dict(material_key=key, status='reserved_exclude' if key in reserved_keys else 'development_identity_candidate',
                fitting_allowed=False)


def independent_video_lengths():
    source = ROOT / 'source_metadata/author_stimulus_database.json'
    database = json.loads(source.read_text(encoding='utf-8'))['database']
    groups = {}
    for name in database:
        match = re.fullmatch(r'(sad|dis|fear|neu|joy|ten|ins)([458])_(\d+)', name)
        if not match:
            raise ValueError('Unexpected named video-segment identity')
        base = match[1] + match[2]
        groups.setdefault(base, []).append(int(match[3]))
    assert len(groups) == 21
    for indices in groups.values():
        assert sorted(indices) == list(range(1, len(indices)+1))
    # These are counts of named segments, independently agreeing with seconds in
    # the current catalogue. No feature arrays, VE8 dummy labels or model outputs
    # are inspected, and no original video content identity is authenticated.
    return {base: len(indices) for base, indices in groups.items()}


def vid_only_identity(body, context):
    """Supplement a missing ordinal vector without inventing an explicit one."""
    from scipy.io import loadmat, whosmat
    schema=whosmat(io.BytesIO(body))
    if len(schema)!=1 or schema[0][0]!='vid' or schema[0][1]!=(1,21) or schema[0][2]!='int32':
        raise ValueError('Not the documented single identity-vector schema; payload unopened')
    values=loadmat(io.BytesIO(body),variable_names=['vid'])['vid'].ravel().tolist()
    lo,hi=(1,21) if context=='ima' else (22,42)
    if len(set(values))!=21 or not all(isinstance(v,int) and lo<=v<=hi for v in values):
        raise ValueError('Invalid single-vector material codes')
    return dict(vid=values, trial_ordinal_supplied=False)


def source_evidence():
    """Publish technical provenance/search summaries, never original source data."""
    specs={
        'author_stimulus_database.json':'https://raw.githubusercontent.com/ncclab-sustech/EmoEEG-MC/d59aed5f9880a19df2a7db71f91caa034b8573c7/tools/annotations/xuxin/xuxin_dataset.json',
        'sciencedb_README.md':'https://china.scidb.cn/download?fileId=e081a9203505e8ebe697bcb8e700d8cc',
        'sciencedb_preprocess.py':'https://china.scidb.cn/download?fileId=5743d9d69cfbe9e73f5f78baadd978c0',
        'historical_video_description.xlsx':'https://s3.amazonaws.com/openneuro.org/ds005540/stimuli/ses-vid/video_description.xlsx?versionId=r9bPBwlAiE0dSa8up9rehb8vqy99ghkD',
    }
    sources=[]
    for name,url in specs.items():
        body=(ROOT/'source_metadata'/name).read_bytes()
        sources.append(dict(name=name,url=url,bytes=len(body),sha256=sha(body)))

    def files(node):
        result=[] if node.get('dir',node.get('type')=='directory') else [node]
        for child in node.get('children') or []:result.extend(files(child))
        return result

    searches=[]
    for name in ('log','experiment','psycho','rating','xlsx','csv'):
        path=ROOT/f'source_metadata/sciencedb_search_{name}.json'
        data=json.loads(path.read_text(encoding='utf-8'))['data']
        matches=[x for x in files(data) if x.get('path','').endswith(('.log','.psyexp','.csv','.xlsx')) or not x.get('dir',False)]
        searches.append(dict(query={'log':'.log','psycho':'.psyexp','xlsx':'.xlsx','csv':'.csv'}.get(name,name),
                             response_sha256=sha(path.read_bytes()),matching_file_entries=len(matches),
                             csv_files_are_only_physiology_features=all('features.csv' in x.get('path','') for x in matches) if name=='csv' else None,
                             original_payloads_retrieved=False))
    checks=json.loads((ROOT/'source_metadata/sciencedb_behaviour_identity_checks.json').read_text())
    cross_archive=[]
    for person in checks:
        cross_archive.append(dict(participant=person['participant'],
                                 tables=[{k:v for k,v in r.items() if k not in ('schema','identities')} for r in person['sources']]))
    return dict(first_party_archive_doi='10.57760/sciencedb.14025',first_party_archive_version='V6',
                archive_page='https://www.scidb.cn/detail?dataSetId=9b182864c9604255a0433d1edae88f0b',
                relevant_sources=sources,archive_metadata_searches=searches,
                cross_archive_behaviour_identity_hash_checks=cross_archive,
                behavioural_comparison_scope='Only trial_number/video_name interpreted; remaining score bytes transported for opaque hashing and discarded, never saved.',
                original_execution_log_found=False,explicit_material_crosswalk_found=False,
                explicit_numbered_score_codebook_found=False,
                rating_values_decoded=0,waveform_samples_decoded=0,models_fitted=0)


def analyze():
    reservation_path = REPO / 'results/development/emo_mc_qualification_2026-10-10/reservation.json'
    reservation = json.loads(reservation_path.read_text())
    archive_text = (ROOT / 'source_metadata/sciencedb_README.md').read_text(encoding='utf-8')
    orders = archive_orders(archive_text)
    original_text = (OLD / 'emo_mc_readme_fixed.response').read_text(encoding='utf-8')
    original_orders = declared_orders(original_text)
    mappings = trigger_mappings(original_text)
    assert all(orders[p] == order for p, order in original_orders.items())
    historical = json.loads((ROOT / 'historical_collection.json').read_text())
    hist = {r['path'].split('/')[0]: r for r in historical}
    identity = json.loads((ROOT / 'identity_collection.json').read_text())
    maps = {(r['participant'], r['context']): r for r in identity}
    supplements=[]
    for context in ('ima','vid'):
        record=maps[('sub-25',context)]
        assert record['status']=='quarantined'
        body=(ROOT/'identity_maps'/context/'sub-25.mat').read_bytes()
        assert sha(body)==record['sha256']
        vectors=vid_only_identity(body,context)
        maps[('sub-25',context)]=dict(record,status='identity_schema_verified_vid_only',identities=vectors)
        supplements.append(dict(participant='sub-25',context=context,source_sha256=sha(body),
                                status='Supplemental single-variable schema validated; original collector quarantine retained in its unchanged output',
                                trial_ordinal_supplied=False,material_codes=21,rating_values_decoded=0))
    write_json(PUBLIC/'identity_schema_supplements.json',dict(supplements=supplements,
                 rating_values_decoded=0,waveform_samples_decoded=0,models_fitted=0))
    behaviour = json.loads((PRIVATE / 'collection.json').read_text())['behaviour']
    behaviours = {b['path'].split('/')[0]: b for b in behaviour}

    # Derive the codebook from the separate author-hosted identity arrays and
    # documented context orders, NOT by zipping EEG trials to behaviour rows.
    lookup = {}
    for context in ('ima', 'vid'):
        source = maps[('sub-37', context)]
        assert source['status'] == 'identity_schema_verified'
        vectors = source['identities']
        assert vectors['trial'] == list(range(1, 22))
        for ordinal, code in zip(vectors['trial'], vectors['vid']):
            lookup[code] = (context, orders['sub-37'][context][ordinal-1])
    assert len(lookup) == 42
    write_json(PUBLIC / 'codebook_candidates.json',
               dict(derivation='Separate ScienceDB V6 trial/vid identity arrays for sub-37 plus its documented per-context order; no EEG/behaviour ordinal zip used to derive these codes.',
                    status='First-party metadata supported; actual execution-log linkage still unverified',
                    entries=[dict(code=k, context=v[0], symbol=v[1]) for k, v in sorted(lookup.items())]))

    summaries = []
    for participant in sorted(hist):
        rows = hist[participant]['events']
        roles = Counter(r['documented_role'] for r in rows)
        blocks = clock_blocks(rows)
        trimmed, initial = separate_initial_state_markers(blocks)
        mapping = mappings.get(participant)
        trials = pair_trials(trimmed, mapping) if mapping and all(k in mapping for k in ('ima','vid','rating','fade')) else []
        counts = Counter(t['context'] for t in trials)
        documented = {c: len(orders[participant][c]) for c in ('ima','vid')} if participant in orders else None
        record = dict(participant=participant, historical_rows=len(rows), documented_marker_roles=dict(roles),
                      backward_clock_jumps=max(int(e['run_path'].rsplit('-',1)[1]) for e in blocks)-1,
                      run_identifiers_supplied=False, discarded_ambiguous_origin_markers=len(initial),
                      candidate_trial_counts=dict(counts), readme_trial_counts=documented,
                      candidate_counts_match_readme=dict(counts)==documented,
                      starts_without_rating_before_next_start=sum(t['rating_start_s'] is None for t in trials),
                      behaviour_rows=len(behaviours[participant]['identities']) if participant in behaviours else None)
        code_comparisons = {}
        archive_comparisons = {}
        b = behaviours.get(participant)
        values = [r.get('video_name_field', r.get('material_id')) for r in b['identities']] if b else []
        coded = len(set(values)) > 1
        parsed = [numeric_identity(v) for v in values] if coded else []
        for context in ('ima','vid'):
            item = maps[(participant,context)]
            if item['status'] not in ('identity_schema_verified','identity_schema_verified_vid_only'):
                archive_comparisons[context] = dict(schema_verified=False)
                continue
            codes = item['identities']['vid']
            symbols = [lookup[k][1] for k in codes]
            expected = orders.get(participant, {}).get(context)
            archive_comparisons[context] = dict(schema_verified=True, planned_identity_entries=len(codes),
                exact_documented_order_match=symbols==expected if expected is not None else None,
                documented_order_is_subsequence=is_subsequence(expected,symbols) if expected is not None else None)
            if coded:
                beh_codes = [k for k in parsed if k in lookup and lookup[k][0]==context]
                code_comparisons[context] = dict(behaviour_codes_equal_archive_identity_order=beh_codes==codes,
                                                codes_in_codebook=all(k in lookup for k in parsed))
        record['archive_identity_order_checks'] = archive_comparisons
        record['behaviour_identity_has_material_codes'] = coded
        record['behaviour_to_archive_order_checks'] = code_comparisons
        record['rating_value_or_execution_link_certified'] = False
        summaries.append(record)

    # Authenticate the old marker exports against the ALREADY authenticated raw
    # annotation channels. This verifies export consistency, not an independent
    # playback-end sensor or an original experiment execution log.
    pilot_checks = []
    for participant, filename in [('sub-37','complete_pilot_events.json'),('sub-60','pilot_events.json')]:
        authenticated = json.loads((PRIVATE / filename).read_text())
        check = compare_events(hist[participant]['events'], authenticated)
        pilot_checks.append(dict(participant=participant, **check))
        assert check['ordered_markers_match_at_edf_precision']

    materials = catalog()
    lengths = independent_video_lengths()
    category_prefix = {'sadness':'sad','disgust':'dis','fear':'fear','neutral':'neu',
                       'joy':'joy','tenderness':'ten','inspiration':'ins'}
    number_map = {}
    video_map = []
    for symbol, count in sorted(lengths.items()):
        matches = [m for m in materials if m['context']=='vid' and symbol.startswith(category_prefix[m['assigned_category']]) and m['duration_s']==count]
        assert len(matches)==1
        material = matches[0]
        number_map[('vid',symbol)] = material['material_key']
        video_map.append(dict(symbol=symbol, current_spreadsheet_number=material['number'],
                              named_segment_count=count, current_catalogue_duration_s=material['duration_s'],
                              status='Independent author-code length consistency; original video content not authenticated'))

    pilot_trials = json.loads((PRIVATE / 'sub-37_trial_identity_pairs.json').read_text())
    rows = behaviours['sub-37']['identities']
    row_for_code = {numeric_identity(r.get('video_name_field',r.get('material_id'))): i+1 for i,r in enumerate(rows)}
    assert len(row_for_code)==42
    counters = Counter()
    manifest = []
    for i, trial in enumerate(pilot_trials):
        context, symbol = trial['context'], trial['readme_symbol']
        ordinal = counters[context]; counters[context]+=1
        code = maps[('sub-37',context)]['identities']['vid'][ordinal]
        assert lookup[code]==(context,symbol)
        rating_row = row_for_code[code]
        mask = protected_mask(context,symbol,number_map,set(reservation['materials']['reserved_keys']))
        record = dict(participant='sub-37', raw_trial_ordinal=i+1, context=context,
                      symbol=symbol, context_trial_ordinal=ordinal+1, archive_identity_code=code,
                      behaviour_row_candidate=rating_row, behaviour_row_equals_raw_trial_ordinal=rating_row==i+1,
                      run_path=trial['run_path'], start_s=trial['start_s'], rating_start_s=trial['rating_start_s'],
                      material_mask=mask, fitting_allowed=False,
                      join_status='Identity metadata triangulated; no original execution-log certificate')
        if context=='vid':
            record['endpoint_checks']=video_endpoint(trial['start_s'],trial['rating_start_s'],lengths[symbol])
        else:
            record['endpoint_checks']=dict(recorded_fade_markers=len(trial['fade_events']),
                       button_or_imagery_end_semantics_verified=False, fitting_allowed=False)
        manifest.append(record)
    assert all(r['behaviour_row_equals_raw_trial_ordinal'] for r in manifest)
    write_json(PUBLIC / 'complete_pilot_join_candidates.json',dict(participant='sub-37',
               trials=manifest, status='Reviewable metadata candidates only; no numerical ratings or EEG samples decoded',
               fitting_allowed=False))

    # Compare possible archive aliases using BOTH context orders. Do not repair
    # participant IDs, override conflicting raw metadata or reallocate anyone.
    planned={}
    for (participant,context),record in maps.items():
        planned.setdefault(participant,{})[context]=[lookup[k][1] for k in record['identities']['vid']]
    override_log=json.loads((ROOT/'source_metadata/fetch_log_12.json').read_text())
    override_proofs=[]
    for record in override_log:
        if 'identity_vectors' not in record:continue
        name=record['name'];context='ima' if '_ima_' in name else 'vid'
        person='sub-11' if '_11.' in name else 'sub-16'
        planned.setdefault(person+'-root-override',dict(planned[person]))[context]=[lookup[k][1] for k in record['identity_vectors']['vid']]
        override_proofs.append({k:v for k,v in record.items() if k!='identity_vectors'})
    aliases=[]
    for person,expected in sorted(orders.items()):
        matches=[dict(archive_identity=s,kind='exact_both_context_orders' if v==expected else 'both_orders_subsequence')
                 for s,v in planned.items() if all(is_subsequence(expected[c],v[c]) for c in ('ima','vid'))]
        if matches!=[dict(archive_identity=person,kind='exact_both_context_orders')]:
            aliases.append(dict(openneuro_participant=person,candidate_archive_identities=matches,
                                unique_candidate=len(matches)==1,verified_crosswalk=False))
    write_json(PUBLIC/'archive_crosswalk_checks.json',dict(nontrivial_or_unresolved_matches=aliases,
                 root_override_metadata=override_proofs,explicit_author_crosswalk_found=False,
                 scope='Matching intended per-context orders only. Equal/similar numeric participant IDs are not assumed to prove identity.',
                 protected_roles_changed=False,fitting_ready=False,rating_values_decoded=0))

    from openpyxl import load_workbook
    older_rows=list(load_workbook(ROOT / 'source_metadata/historical_video_description.xlsx',read_only=True,data_only=True).active.values)[1:]
    new_rows=list(load_workbook(PRIVATE / 'discovery/video_description.xlsx',read_only=True,data_only=True).active.values)[1:]
    assert len(older_rows)==len(new_rows)==21
    duration_changes=[dict(number=o[0], older_duration_s=o[4], current_duration_s=n[4])
                      for o,n in zip(older_rows,new_rows) if o[4]!=n[4]]
    assert all(o[0]==n[0] and o[3]==n[3] for o,n in zip(older_rows,new_rows))
    identities_inconsistent=[s['participant'] for s in summaries
        if any(v.get('exact_documented_order_match') is False for v in s['archive_identity_order_checks'].values())]
    not_subsequence=[s['participant'] for s in summaries
        if any(v.get('documented_order_is_subsequence') is False for v in s['archive_identity_order_checks'].values())]
    results=dict(historical_event_tables=len(historical), historical_markers=sum(len(r['events']) for r in historical),
                 historical_durations_all_zero=True, historical_run_ids_supplied=False,
                 historical_export_pilot_checks=pilot_checks,
                 independent_identity_schema_verified=sum(r['status']=='identity_schema_verified' for r in identity),
                 independent_identity_schema_quarantines=[dict(participant=r['participant'],context=r['context'],reason=r['error']) for r in identity if r['status']=='quarantined'],
                 supplemental_identity_only_files=2,total_identity_only_files_inspected=120,
                 participant_checks=summaries,
                 archive_orders_not_exact=identities_inconsistent, archive_orders_not_subsequence=not_subsequence,
                 sub54_metadata_order_recovered=True, sub54_original_quarantine_role_retained=True,
                 video_symbol_number_consistency=video_map, older_catalogue_duration_changes=duration_changes,
                 imagery_symbol_to_spreadsheet_number_verified=False,
                 complete_pilot_identity_candidate_rows=len(manifest), complete_pilot_row_order_checks=42,
                 video_endpoint_mismatches_over_1s=[dict(symbol=r['symbol'],run_path=r['run_path'],**r['endpoint_checks']) for r in manifest if r['context']=='vid' and abs(r['endpoint_checks']['candidate_end_difference_s'])>1],
                 actual_video_end_or_button_semantics_verified=False,
                 rating_column_semantics_explicitly_numbered=False,
                 reservation_sha256=sha(reservation_path.read_bytes()), protected_roles_changed=False,
                 no_reallocation=True, rating_values_decoded=0,waveform_samples_decoded=0,models_fitted=0,
                 research_question_changed=False, fitting_ready=False,
                 analysis_script_sha256=sha(Path(__file__).read_bytes()),
                 open_gates=['Actual execution-to-rating-row linkage for unresolved/constant-code/incomplete cases',
                             'Imagery symbol-to-number map for the frozen protected material mask',
                             'Actual video end and imagery button/fade semantics; no cross-rating stitching',
                             'Explicit numbered score-column semantics and annotation-noise/run reconciliation'])
    write_json(PUBLIC / 'join_resolution.json',results)
    write_json(PUBLIC/'source_discovery.json',source_evidence())
    print(json.dumps({k:results[k] for k in ['historical_event_tables','historical_markers','independent_identity_schema_verified','archive_orders_not_exact','archive_orders_not_subsequence','sub54_metadata_order_recovered','complete_pilot_identity_candidate_rows','video_endpoint_mismatches_over_1s','fitting_ready']},indent=2))


if __name__ == '__main__':
    analyze()
